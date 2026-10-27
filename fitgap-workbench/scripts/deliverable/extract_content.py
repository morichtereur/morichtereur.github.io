#!/usr/bin/env python3
"""Reads a finished fit-gap deliverable (both tabs + Cover) back into the content JSON that build_fitgap.py consumes.

    python scripts/deliverable/extract_content.py <deliverable.xlsx> <content.json>

Used to regenerate a reference file from its template (format regression test) and to compare deliverables.
"""
import json
import re
import sys

import openpyxl

from fitgap_layout import BLOCK_FIELDS, TemplateLayout


def read_tab(ws, codes):
    lay = TemplateLayout(ws)
    rows = {}
    seen = {}
    for (r, lv, pid, name, desc) in lay.rows():
        key = pid or f'row:{r}'
        if pid:
            seen[pid] = seen.get(pid, 0) + 1
            if seen[pid] > 1:
                key = f'{pid}#{seen[pid]}'
        per = {}
        yellow = {}
        for b, code in enumerate(codes):
            vals = {}
            cols = lay.blocks[b]['cols']
            for k in BLOCK_FIELDS:
                c = cols.get(k)
                v = ws.cell(r, c).value if c else None
                vals[k] = None if v in (None, '', '\xa0') else v
                if c and ws.cell(r, c).fill.fill_type == 'solid' and str(ws.cell(r, c).fill.fgColor.rgb).endswith('FFFF00'):
                    yellow.setdefault(code, []).append(k)
            per[code] = vals
        rows[key] = (per, yellow)
    return lay, rows


def read_cover(ws):
    sections = []
    for r in range(1, ws.max_row + 1):
        b, c = ws.cell(r, 2), ws.cell(r, 3)
        if b.value in (None, ''):
            continue
        if b.font.b and (b.font.sz or 0) >= 11 and c.value in (None, ''):
            sections.append({'title': str(b.value), 'items': []})
        elif sections:
            sections[-1]['items'].append([str(b.value), '' if c.value is None else str(c.value)])
    return sections


def extract(path, out):
    wb = openpyxl.load_workbook(path)
    fin = next(n for n in wb.sheetnames if 'fit-gap' in n.lower() and not n.lower().endswith('_vws'))
    vws = next((n for n in wb.sheetnames if n.lower().endswith('_vws')), None)
    lay = TemplateLayout(wb[fin])
    labels = [b['label'] for b in lay.blocks]
    codes = [(re.search(r'\((\w+)\)\s*$', l) or re.search(r'(\w+)$', l)).group(1) for l in labels]
    _, final_rows = read_tab(wb[fin], codes)
    _, vws_rows = read_tab(wb[vws], codes) if vws else (None, {})
    m = re.match(r'(.*?)_([A-Z]+)$', fin.replace('Fit-Gap_', ''))
    content = {
        'ws': lay.ws.cell(lay.hdr + 2, 3).value.split()[0] if lay.ws.cell(lay.hdr + 2, 3).value else '',
        'final_sheet': fin, 'vws_sheet': vws or fin + '_vWS',
        'entities': [{'code': c, 'label': l, 'template_block': i} for i, (c, l) in enumerate(zip(codes, labels))],
        'cover': read_cover(wb['Cover']) if 'Cover' in wb.sheetnames else [],
        'rows': {},
    }
    for key, (per, yellow) in final_rows.items():
        spec = {'final': per}
        if key in vws_rows and vws_rows[key][0] != per:
            spec['vws'] = vws_rows[key][0]
        if yellow:
            spec['yellow'] = yellow
        content['rows'][key] = spec
    json.dump(content, open(out, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print(f'{path}: {len(content["rows"])} rows, entities {codes}, cover sections {len(content["cover"])} → {out}')


if __name__ == '__main__':
    extract(sys.argv[1], sys.argv[2])
