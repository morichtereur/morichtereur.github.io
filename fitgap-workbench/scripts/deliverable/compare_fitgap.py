#!/usr/bin/env python3
"""Compares two fit-gap deliverables cell by cell (values, styles, merges, widths, heights, CF, DV, panes, media).

    python scripts/deliverable/compare_fitgap.py <built.xlsx> <reference.xlsx> [--sheet <name>] [--max 15]

Format regression test: rebuild a reference from its template + extracted content and compare.
"""
import argparse
import zipfile
from collections import Counter

import openpyxl
from openpyxl.utils import get_column_letter as L


def color(c):
    if c is None:
        return '-'
    if c.type == 'rgb':
        return str(c.rgb)[-6:]
    if c.type == 'theme':
        return f't{c.theme}{"%+.2f" % c.tint if c.tint else ""}'
    return c.type


def style(c):
    f, fl, a, b = c.font, c.fill, c.alignment, c.border
    return {'font': (f.name, f.sz, bool(f.b), color(f.color)), 'fill': color(fl.fgColor) if fl.fill_type else '-',
            'align': (a.horizontal, a.vertical, bool(a.wrap_text)),
            'border': tuple(((getattr(b, s).style if getattr(b, s) else None), color(getattr(b, s).color) if getattr(b, s) and getattr(b, s).color else '-') for s in ('left', 'right', 'top', 'bottom'))}


def norm(v):
    v = v.strip() if isinstance(v, str) else v
    return None if v == '' else v


def compare(built, ref, sheet=None, limit=15):
    wa, wb = openpyxl.load_workbook(built), openpyxl.load_workbook(ref)
    print('sheets built :', wa.sheetnames, '| active', wa.active.title)
    print('sheets ref   :', wb.sheetnames, '| active', wb.active.title)
    za, zb = zipfile.ZipFile(built), zipfile.ZipFile(ref)
    media = lambda z: sorted(n for n in z.namelist() if n.startswith('xl/media/'))
    print('media files  :', len(media(za)), 'vs', len(media(zb)))
    names = [sheet] if sheet else [n for n in wb.sheetnames if 'fit-gap' in n.lower()]
    for n in names:
        if n not in wa.sheetnames:
            print(f'!! sheet {n} missing in built'); continue
        a, b = wa[n], wb[n]
        print(f'\n=== {n}: dims {a.dimensions} vs {b.dimensions}; freeze {a.freeze_panes} vs {b.freeze_panes}; filter {a.auto_filter.ref} vs {b.auto_filter.ref}; gridlines {a.sheet_view.showGridLines} vs {b.sheet_view.showGridLines}')
        ma, mb = set(map(str, a.merged_cells.ranges)), set(map(str, b.merged_cells.ranges))
        print('merges: only built', sorted(ma - mb)[:limit], '| only ref', sorted(mb - ma)[:limit])
        wdiff = [(L(c), a.column_dimensions[L(c)].width, b.column_dimensions[L(c)].width) for c in range(1, max(a.max_column, b.max_column) + 1)
                 if abs((a.column_dimensions[L(c)].width or 0) - (b.column_dimensions[L(c)].width or 0)) > 0.3]
        print(f'width diffs ({len(wdiff)}):', wdiff[:limit])
        hd = [(r, a.row_dimensions[r].height, b.row_dimensions[r].height) for r in range(1, max(a.max_row, b.max_row) + 1)
              if abs((a.row_dimensions[r].height or 15) - (b.row_dimensions[r].height or 15)) > 0.6]
        print(f'row height diffs ({len(hd)} of {max(a.max_row, b.max_row)}):', hd[:limit])
        vdiff, sdiff = [], Counter()
        sdiff_ex = {}
        for r in range(1, max(a.max_row, b.max_row) + 1):
            for c in range(1, max(a.max_column, b.max_column) + 1):
                ca, cb = a.cell(r, c), b.cell(r, c)
                if norm(ca.value) != norm(cb.value):
                    vdiff.append((ca.coordinate, str(ca.value)[:40], str(cb.value)[:40]))
                sa, sb = style(ca), style(cb)
                for k in sa:
                    if sa[k] != sb[k]:
                        key = (k, L(c)) if r > 10 else (k, f'R{r}')
                        sdiff[key] += 1
                        sdiff_ex.setdefault(key, (ca.coordinate, sa[k], sb[k]))
        print(f'value diffs ({len(vdiff)}):')
        for x in vdiff[:limit]:
            print('   ', x)
        print(f'style diffs ({sum(sdiff.values())} cells; by kind/column):')
        for key, nn in sdiff.most_common(limit):
            print('   ', key, nn, sdiff_ex[key])
        cfa = sorted((str(k.sqref), ru.type, str(ru.formula)) for k, rules in a.conditional_formatting._cf_rules.items() for ru in rules)
        cfb = sorted((str(k.sqref), ru.type, str(ru.formula)) for k, rules in b.conditional_formatting._cf_rules.items() for ru in rules)
        print('CF equal:', cfa == cfb, '' if cfa == cfb else (cfa[:4], cfb[:4]))
        dva = sorted((str(d.sqref), d.formula1) for d in a.data_validations.dataValidation)
        dvb = sorted((str(d.sqref), d.formula1) for d in b.data_validations.dataValidation)
        print('DV equal:', dva == dvb, '' if dva == dvb else (dva, dvb))
    if 'Cover' in wa.sheetnames and 'Cover' in wb.sheetnames:
        a, b = wa['Cover'], wb['Cover']
        d = [(a.cell(r, c).coordinate, str(a.cell(r, c).value)[:30], str(b.cell(r, c).value)[:30]) for r in range(1, max(a.max_row, b.max_row) + 1) for c in (2, 3) if norm(a.cell(r, c).value) != norm(b.cell(r, c).value)]
        print(f'\nCover value diffs ({len(d)}):', d[:limit], '| widths', a.column_dimensions['B'].width, a.column_dimensions['C'].width, 'vs', b.column_dimensions['B'].width, b.column_dimensions['C'].width)


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('built'); ap.add_argument('ref'); ap.add_argument('--sheet'); ap.add_argument('--max', type=int, default=15)
    a = ap.parse_args()
    compare(a.built, a.ref, a.sheet, a.max)
