#!/usr/bin/env python3
"""Builds the country fit-gap deliverable from a capture-sheet template and a resolved content file.

    python scripts/deliverable/build_fitgap.py --template <capture.xlsx> --content <content.json> --out <file.xlsx>

Keeps every other sheet of the template (Cover, >>, Process Flows, Process Capture …, images included) and replaces
the single fit-gap sheet by two tabs, "<WS>_<COUNTRY>_vWS" (full session notes) and "<WS>_<COUNTRY>" (final, active):
taxonomy B–E | standard split F/G, system H, initiatives I | spacer | 17-column block per entity | spacer | the
template's remaining columns (process information, automation, change, RACI). Column roles are read from the header
texts (fitgap_layout.py), so the same code serves every stream template.

Content file (JSON, produced per stream from the capture file and the session notes – see extract_content.py for
the shape): entities, cover sections and per-row values for both tabs; fields missing in the content fall back to
the template's pre-assessment values of the entity's block, an explicit null clears the cell.
"""
import argparse
import json
import math
import re
from copy import copy

import openpyxl
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

from fitgap_layout import BLOCK_FIELDS, RATINGS, TemplateLayout

# ---- canonical header texts (reference files), template typos corrected
TYPO = {'Inititaves': 'Initiatives', 'initiatve': 'initiative', 'activitiy': 'activity', 'additonal': 'additional',
        'addiitonal': 'additional', 'requiered': 'required', 'Aseesmnent': 'Assessment', 'ASSSSMENT': 'ASSESSMENT',
        'contitions': 'conditions', 'performes': 'performs', 'Standardizaion': 'Standardization', 'activites': 'activities',
        'Assesment': 'Assessment', 'Managment': 'Management', 'recieve': 'receive', 'seperate': 'separate'}
GROUP_TITLES = {'standard': 'TO BE PROCESS SPECIFIC INFORMATION', 'asis': 'PRE-FILLED AS-IS PROCESS SPECIFICATION ',
                'fitgap': 'FIT-GAP ASSESSMENT', 'detail': 'DETAILED GAP ASSESSMENT'}
SUB = {'std_le': 'Activity Split', 'std_sys': 'System & Tool ', 'std_init': 'Initiatives', 'as_le': 'Activity Split As-is ',
       'sys': 'System & Tool ', 'gap': 'GAP', 'info': 'Comments', 'tb_local': 'Activity Split', 'fte': 'FTE ', 'cmt': 'Comments',
       'd_sys': 'IT Systems', 'd_cur': 'Local Business Requirements', 'd_legal': 'Legal '}
NAMES = {'level': 'Process Level', 'pid': 'Process ID', 'name': 'Process ', 'desc': 'Activity description',
         'std_le': 'LE/ Country', 'std_gbs': 'GBS', 'std_sys': 'Used Systems', 'std_init': 'Planned initiative',
         'as_le': 'LE/ Country / Other', 'as_gbs': 'GBS', 'sys': 'Used Systems', 'gap': 'GAP', 'info': 'Additional Information',
         'tb_local': 'Local Accounting', 'tb_other': 'Other', 'tb_gbs': 'GBS', 'fte': 'FTE ', 'cmt': 'Comment',
         'd_sys': 'System & Applications', 'd_wf': 'Workflow Tools', 'd_cur': 'Current Process', 'd_ext': 'External Party',
         'd_lang': 'Language requirements', 'd_legal': 'Legal / Regulatory Requirements', 'd_ics': 'Internal Controls'}
DESC = {'level': 'Process Level', 'pid': 'Process ID', 'name': 'Process ', 'desc': 'Activity description',
        'std_le': 'The local LE/country performs this activity', 'std_gbs': 'GBS performs activity ',
        'std_sys': 'Platform/app activity is executed', 'std_init': 'What initiative is planned per process?',
        'as_le': 'The LE/country/others perform this activity', 'as_gbs': 'GBS performs activity ',
        'sys': 'Platform/app activity is executed', 'gap': 'GAP Assessment',
        'info': 'Any additional findings / local requirements etc.', 'tb_local': 'The LE/country perform this activity',
        'tb_other': 'Other perform this activity', 'tb_gbs': 'GBS performs activity ',
        'fte': 'Additional information about FTE of LE per process steps to perform the activity',
        'cmt': 'Any additional findings / local requirements etc.',
        'd_sys': 'What IT systems are used to perform this activity?',
        'd_wf': 'What workflow tools are used to perform this activity?',
        'd_cur': 'Are there any country or location specifics?',
        'd_ext': 'Are there any activities outsourced to a local service provider? ',
        'd_lang': 'Are there any languages required (other than English) to perform this activity?',
        'd_legal': 'Are there any local legal requirements or regulations that need to be adhered to?',
        'd_ics': 'Are there any special local legal control that are implemented? '}
WIDTH = {'A': 4.0, 'level': 7.5, 'pid': 16.5, 'name': 51.5, 'desc': 35.4, 'std_le': 21.5, 'std_gbs': 16.5, 'std_sys': 23.9,
         'std_init': 22.0, 'spacer': 2.0, 'as_le': 12.0, 'as_gbs': 8.0, 'sys': 14.0, 'gap': 20.0, 'info': 26.0,
         'tb_local': 10.0, 'tb_other': 8.0, 'tb_gbs': 8.0, 'fte': 30.0, 'cmt': 70.0, 'd_sys': 18.0, 'd_wf': 16.0,
         'd_cur': 26.0, 'd_ext': 16.0, 'd_lang': 12.0, 'd_legal': 18.0, 'd_ics': 18.0}
CENTER = {'level', 'pid', 'std_le', 'std_gbs', 'as_le', 'as_gbs', 'gap', 'tb_local', 'tb_other', 'tb_gbs'}
LEFT_MID = {'name', 'desc', 'sys'}
CF_STYLE = {'NO GAP': ('C6EFCE', '006100'), 'HARMONIZATION GAP': ('FFC7CE', '9C0006'), 'STANDARDIZATION GAP': ('FFEB9C', '9C5700'),
            'BOTH GAPS': ('FFC7CE', '9C0006'), 'N/A': (None, '7F7F7F'), 'OUT OF SCOPE': (None, '7F7F7F')}
LEVEL_STYLE = {1: ('1F3864', 'FFFFFF'), 2: ('D9E1F2', None), 3: ('F2F2F2', None)}
GREY = Side(style='thin', color='BFBFBF')
BORDER = Border(left=GREY, right=GREY, top=GREY, bottom=GREY)
YELLOW = PatternFill('solid', start_color='FFFF00', end_color='FFFF00')
# Row height: Verdana 8 wrapped text, calibrated on the reference files (chars per width unit, points per line).
CHARS_PER_UNIT, LINE_PT, PAD_PT, MIN_PT = 0.9, 11.5, 3.0, 14.5


def theme(t, tint=0.0):
    return PatternFill('solid', fgColor=Color(theme=t, tint=tint))


def fix_typos(s):
    if not isinstance(s, str):
        return s
    for a, b in TYPO.items():
        s = s.replace(a, b)
    return s


HEADER_OVERRIDE = {'gap initial assessment': 'GAP', 'gap aseesmnent': 'GAP Assessment', 'gap assessment': 'GAP Assessment',
                   'initial gap assessment': 'GAP Assessment', 'gap': 'GAP', 'comments': 'Comments', 'comment': 'Comment'}


def header_text(src, row, tcol, role, table):
    """Header text for a standard/block column: the template's own text (typos fixed, line breaks joined) when the
    template has the column, else the canonical text of the reference files."""
    v = src.cell(row, tcol).value if tcol else None
    if v in (None, '', '\xa0'):
        return table[role]
    t = re.sub(r'\s*\n\s*', ' ', fix_typos(str(v))).replace('  ', ' ')
    key = t.strip().lower()
    if role in ('gap', 'info', 'cmt') and key in HEADER_OVERRIDE:
        return HEADER_OVERRIDE[key] if role == 'gap' or table is NAMES else table[role]
    return t


def copy_style(src, dst):
    dst.font, dst.fill, dst.border, dst.alignment, dst.number_format = copy(src.font), copy(src.fill), copy(src.border), copy(src.alignment), src.number_format


def shift_formula(formula, colmap):
    """Re-points same-sheet A1 references of a copied formula to the new column positions (rows are unchanged)."""
    def sub(m):
        if m.group(1) or m.group(0).startswith('!'):
            return m.group(0)
        col = m.group(3)
        from openpyxl.utils import column_index_from_string as C
        try:
            new = colmap.get(C(col))
        except ValueError:
            new = None
        return m.group(0) if new is None else f'{m.group(2)}{L(new)}{m.group(4)}{m.group(5)}'
    return re.sub(r"(['!])?(\$?)([A-Z]{1,3})(\$?)(\d+)(?![\w(])", sub, formula)


def text_lines(text, width):
    if text in (None, ''):
        return 0
    cpl = max(1.0, width * CHARS_PER_UNIT)
    return sum(max(1, math.ceil(len(par) / cpl)) for par in str(text).split('\n'))


class Plan:
    """Output column plan: list of (role, template_col or None, group) in sheet order."""

    def __init__(self, layout, n_entities):
        self.layout = layout
        self.cols = [('A', None, None)]
        for k in ('level', 'pid', 'name', 'desc'):
            self.cols.append((k, layout.taxonomy[k], 'taxonomy'))
        for k in ('std_le', 'std_gbs', 'std_sys', 'std_init'):
            if layout.std.get(k) or k != 'std_init':
                self.cols.append((k, layout.std.get(k), 'standard'))
        self.blocks = []
        for b in range(n_entities):
            self.cols.append(('spacer', None, None))
            start = len(self.cols)
            tb = layout.blocks[min(b, len(layout.blocks) - 1)]['cols']
            for k in BLOCK_FIELDS:
                self.cols.append((k, tb.get(k), f'block{b}'))
            self.blocks.append({k: start + i + 1 for i, k in enumerate(BLOCK_FIELDS)})
        self.cols.append(('spacer', None, None))
        self.trailing_start = len(self.cols) + 1
        for g in layout.trailing:
            for c in range(g['start'], g['end'] + 1):
                self.cols.append(('trailing', c, g['kind']))
        self.index = {}
        for i, (role, _, grp) in enumerate(self.cols, start=1):
            self.index.setdefault((grp, role), i)
        self.n = len(self.cols)

    def col(self, grp, role):
        return self.index[(grp, role)]


def build_sheet(wb, layout, plan, content, variant, title, index):
    ws = wb.create_sheet(title, index)
    src = layout.ws
    hdr = layout.hdr
    ents = content['entities']
    # ---- column widths
    for i, (role, tcol, grp) in enumerate(plan.cols, start=1):
        if role == 'trailing':
            w = src.column_dimensions[L(tcol)].width
            ws.column_dimensions[L(i)].width = w if w else 13.0
        else:
            ws.column_dimensions[L(i)].width = WIDTH.get(role, 13.0)
    # ---- header rows: B column texts and trailing columns copied from the template, the rest built
    for r in range(1, hdr + 1):
        ws.row_dimensions[r].height = src.row_dimensions[r].height
        for i, (role, tcol, grp) in enumerate(plan.cols, start=1):
            if role == 'trailing' or role == 'A' or (grp == 'taxonomy' and r <= layout.group_row - 1):
                sc = src.cell(r, tcol if role != 'A' else 1)
                dc = ws.cell(r, i)
                if sc.value not in (None, '\xa0'):
                    dc.value = fix_typos(sc.value)
                copy_style(sc, dc)
    # trailing header merges, remapped
    shift = {tcol: i for i, (role, tcol, grp) in enumerate(plan.cols, start=1) if role == 'trailing'}
    for m in src.merged_cells.ranges:
        if m.max_row <= hdr and m.min_col in shift and m.max_col in shift:
            ws.merge_cells(start_row=m.min_row, start_column=shift[m.min_col], end_row=m.max_row, end_column=shift[m.max_col])
    g_row, s_row, n_row = layout.group_row, layout.sub_row, layout.names_row
    bold = lambda sz, color=None: Font(name='Verdana', size=sz, bold=True, color=color)
    white = Color(theme=0)
    # taxonomy titles
    c_level, c_desc = plan.col('taxonomy', 'level'), plan.col('taxonomy', 'desc')
    ws.cell(s_row, c_level).value = layout.taxonomy_title
    ws.cell(s_row, c_level).font = bold(11, Color(theme=1)); ws.cell(s_row, c_level).fill = theme(0, -0.15)
    ws.cell(s_row, c_level).alignment = Alignment(horizontal='left', vertical='center')
    ws.merge_cells(start_row=s_row, start_column=c_level, end_row=s_row, end_column=c_desc)
    ws.cell(n_row, c_level).value = layout.steps_title
    ws.cell(n_row, c_level).font = bold(10, white); ws.cell(n_row, c_level).fill = theme(4)
    ws.cell(n_row, c_level).alignment = Alignment(horizontal='left', vertical='center')
    ws.merge_cells(start_row=n_row, start_column=c_level, end_row=n_row, end_column=c_desc)
    for k in ('level', 'pid', 'name', 'desc'):
        c = ws.cell(hdr, plan.col('taxonomy', k)); c.value = DESC[k]
        c.font = Font(name='Verdana', size=8); c.fill = PatternFill('solid', start_color='F2F2F2', end_color='F2F2F2')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); c.border = BORDER
    # standard split
    std_cols = [i for i, (role, _, grp) in enumerate(plan.cols, start=1) if grp == 'standard']
    c = ws.cell(g_row, std_cols[0]); c.value = GROUP_TITLES['standard']; c.font = bold(10, white); c.fill = theme(3, 0.5)
    c.alignment = Alignment(horizontal='left', vertical='center')
    ws.merge_cells(start_row=g_row, start_column=std_cols[0], end_row=g_row, end_column=std_cols[-1])
    for i in std_cols:
        role, tcol = plan.cols[i - 1][0], plan.cols[i - 1][1]
        if role in SUB:
            c = ws.cell(s_row, i); c.value = header_text(src, s_row, tcol, role, SUB); c.font = bold(8); c.fill = theme(3, 0.9)
            c.alignment = Alignment(horizontal='center', vertical='center')
        c = ws.cell(n_row, i); c.value = header_text(src, n_row, tcol, role, NAMES); c.font = bold(9, white); c.fill = theme(3, 0.5)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c = ws.cell(hdr, i); c.value = header_text(src, hdr, tcol, role, DESC); c.font = Font(name='Verdana', size=8)
        c.fill = PatternFill('solid', start_color='F2F2F2', end_color='F2F2F2'); c.border = BORDER
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.merge_cells(start_row=s_row, start_column=std_cols[0], end_row=s_row, end_column=std_cols[1])
    # entity blocks
    for b, ent in enumerate(ents):
        bc = plan.blocks[b]
        for k, span, fill in (('as_le', 5, theme(2, -0.25)), ('tb_local', 5, theme(3, 0.1)), ('d_sys', 7, theme(3, 0.1))):
            c = ws.cell(g_row, bc[k]); c.value = GROUP_TITLES[{'as_le': 'asis', 'tb_local': 'fitgap', 'd_sys': 'detail'}[k]]
            c.font = bold(10, white); c.fill = fill; c.alignment = Alignment(horizontal='left', vertical='center')
            ws.merge_cells(start_row=g_row, start_column=bc[k], end_row=g_row, end_column=bc[k] + span - 1)
        for k in ('as_le', 'tb_local'):
            c = ws.cell(layout.label_row or g_row - 2, bc[k]); c.value = ent['label']; c.font = bold(10); c.alignment = Alignment(vertical='center')
        tcols = layout.blocks[min(ent.get('template_block', b), len(layout.blocks) - 1)]['cols']
        for k, span in (('as_le', 2), ('sys', 1), ('gap', 1), ('info', 1), ('tb_local', 3), ('fte', 1), ('cmt', 1),
                        ('d_sys', 2), ('d_cur', 3), ('d_legal', 2)):
            c = ws.cell(s_row, bc[k]); c.value = header_text(src, s_row, tcols.get(k), k, SUB)
            if k in ('as_le', 'sys', 'info'):
                c.font = bold(8); c.fill = theme(2)
            elif k == 'gap':
                c.font = bold(8); c.fill = theme(0)
            elif k in ('tb_local', 'fte', 'cmt'):
                c.font = bold(8); c.fill = theme(3, 0.9)
            else:
                c.font = bold(8, white); c.fill = theme(3, 0.5)
            c.alignment = Alignment(horizontal='center', vertical='center')
            if span > 1:
                ws.merge_cells(start_row=s_row, start_column=bc[k], end_row=s_row, end_column=bc[k] + span - 1)
        for k in BLOCK_FIELDS:
            c = ws.cell(n_row, bc[k]); c.value = header_text(src, n_row, tcols.get(k), k, NAMES); c.font = bold(8, white)
            c.fill = theme(2, -0.25) if k in ('as_le', 'as_gbs', 'sys', 'info') else theme(3, 0.1) if k in ('gap', 'tb_local', 'tb_other', 'tb_gbs', 'fte', 'cmt') else theme(3, 0.75)
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); c.border = BORDER
            c = ws.cell(hdr, bc[k]); c.value = header_text(src, hdr, tcols.get(k), k, DESC); c.border = BORDER
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            if k in ('as_le', 'as_gbs', 'sys', 'info'):
                c.font = Font(name='Verdana', size=8); c.fill = PatternFill('solid', start_color='F2F2F2', end_color='F2F2F2')
            else:
                c.font = Font(name='Verdana', size=8, color='000000'); c.fill = theme(0, -0.05)
    # ---- data rows
    rows = content['rows']
    seen = {}
    out_r = hdr
    for (r, lv, pid, name, desc) in layout.rows():
        out_r += 1
        key = pid or f'row:{r}'
        if pid:
            seen[pid] = seen.get(pid, 0) + 1
            if seen[pid] > 1:
                key = f'{pid}#{seen[pid]}'
        spec = rows.get(key, {})
        values = {}
        for i, (role, tcol, grp) in enumerate(plan.cols, start=1):
            if role in ('A', 'spacer'):
                continue
            if grp == 'taxonomy':
                v = src.cell(r, tcol).value
            elif grp == 'standard':
                v = src.cell(r, tcol).value if tcol else None
            elif role == 'trailing':
                v = src.cell(r, tcol).value
                if isinstance(v, str) and v.startswith('='):
                    v = shift_formula(v, shift)
            else:
                b = int(grp[5:]); ent = ents[b]
                base = layout.block_values(r, ent.get('template_block', b))
                var = spec.get(variant) or spec.get('final') or {}
                ov = var.get(ent['code'], {})
                v = ov[role] if role in ov else base.get(role)
            if isinstance(v, str):
                v = v.replace('\xa0', ' ')
                v = v if v.strip() else None
            values[i] = v
            ws.cell(out_r, i).value = v
        # styles
        fill_hex, font_color = LEVEL_STYLE.get(lv, (None, None))
        is_bold = lv in LEVEL_STYLE
        yellow = spec.get('yellow', {})
        for i, (role, tcol, grp) in enumerate(plan.cols, start=1):
            c = ws.cell(out_r, i)
            if role == 'A':
                c.font = Font(name='Verdana', size=8); c.alignment = Alignment(vertical='center')
                continue
            c.font = Font(name='Verdana', size=8, bold=is_bold, color=font_color)
            if role == 'spacer':
                c.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
                continue
            c.border = BORDER
            if fill_hex:
                c.fill = PatternFill('solid', start_color=fill_hex, end_color=fill_hex)
            if grp and grp.startswith('block') and role in yellow.get(ents[int(grp[5:])]['code'], []):
                c.fill = YELLOW
            if role in CENTER:
                c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            elif role in LEFT_MID:
                c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
            else:
                c.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
        # row height from the longest wrapped text
        lines = 1
        for i, v in values.items():
            if v is None:
                continue
            lines = max(lines, text_lines(v, ws.column_dimensions[L(i)].width or 13.0))
        ws.row_dimensions[out_r].height = max(MIN_PT, lines * LINE_PT + PAD_PT) if lines > 1 else MIN_PT
    last = out_r
    # ---- conditional formatting, validation, panes, filter, view
    first = hdr + 1
    dv_gap = DataValidation(type='list', formula1='"' + ','.join(RATINGS) + '"', allow_blank=True)
    dv_x = DataValidation(type='list', formula1='"x"', allow_blank=True)
    for bc in plan.blocks:
        rng = f'{L(bc["gap"])}{first}:{L(bc["gap"])}{last}'
        for rating, (bg, fg) in CF_STYLE.items():
            kw = {'font': Font(color=fg)}
            if bg:
                kw['fill'] = PatternFill('solid', start_color=bg, end_color=bg)
            ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=[f'"{rating}"'], **kw))
        dv_gap.add(rng)
        for k in ('tb_local', 'tb_other', 'tb_gbs'):
            dv_x.add(f'{L(bc[k])}{first}:{L(bc[k])}{last}')
    ws.add_data_validation(dv_gap); ws.add_data_validation(dv_x)
    ws.freeze_panes = f'{L(plan.col("standard", "std_le"))}{first}'
    ws.auto_filter.ref = f'A{hdr}:{L(plan.n)}{last}'
    ws.sheet_view.showGridLines = False
    return ws


def write_cover(wb, content):
    sections = content.get('cover')
    if not sections:
        return
    ws = wb['Cover'] if 'Cover' in wb.sheetnames else wb.create_sheet('Cover', 0)
    for row in ws.iter_rows():
        for c in row:
            c.value = None; c.font = Font(name='Arial', size=9); c.fill = PatternFill(); c.border = Border()
    for m in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(m))
    ws.column_dimensions['B'].width = content.get('cover_key_width', 30.0)
    ws.column_dimensions['C'].width = 110.0
    r = 3
    for sec in sections:
        ws.cell(r, 2).value = sec['title']; ws.cell(r, 2).font = Font(name='Arial', size=11, bold=True)
        r += 1
        for key, text in sec['items']:
            ws.cell(r, 2).value = key; ws.cell(r, 2).font = Font(name='Arial', size=9, bold=True)
            ws.cell(r, 2).alignment = Alignment(vertical='top')
            ws.cell(r, 3).value = text; ws.cell(r, 3).font = Font(name='Arial', size=9)
            ws.cell(r, 3).alignment = Alignment(wrap_text=True, vertical='top')
            ws.row_dimensions[r].height = max(12.75, text_lines(text, 110.0 * 1.0) * 12.0 + 1.5) if text and len(str(text)) > 118 else None
            r += 1
        r += 1
    ws.sheet_view.showGridLines = False


def build(template, content_path, out):
    content = json.load(open(content_path, encoding='utf-8'))
    wb = openpyxl.load_workbook(template)
    src_name = content.get('template_sheet') or next(n for n in wb.sheetnames if 'fit' in n.lower() and 'gap' in n.lower())
    wb[src_name].title = '_template'
    src_name = '_template'
    layout = TemplateLayout(wb[src_name])
    print('template:', layout.describe())
    plan = Plan(layout, len(content['entities']))
    idx = wb.sheetnames.index(src_name)
    build_sheet(wb, layout, plan, content, 'vws', content['vws_sheet'], idx)
    final = build_sheet(wb, layout, plan, content, 'final', content['final_sheet'], idx + 1)
    del wb[src_name]
    for n in content.get('drop_sheets', []):
        if n in wb.sheetnames:
            del wb[n]
    write_cover(wb, content)
    wb.active = wb.sheetnames.index(final.title)
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws.title == final.title
    wb.save(out)
    print(f'written {out}: {wb.sheetnames}, final tab {final.title} rows {layout.hdr + 1}–{final.max_row}, {plan.n} columns')
    return out


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--template', required=True); ap.add_argument('--content', required=True); ap.add_argument('--out', required=True)
    a = ap.parse_args()
    build(a.template, a.content, a.out)
