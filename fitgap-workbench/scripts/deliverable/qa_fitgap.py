#!/usr/bin/env python3
"""Read-only quality checks for a fit-gap deliverable, as a markdown report.

    python scripts/deliverable/qa_fitgap.py <deliverable.xlsx> [--capture <capture.xlsx>] [--people "Name,Name"] [--out report.md]

Checks: structure (sheets, tabs, columns, panes, filter, CF/DV, static fills on rating cells), ratings vs capture,
content rules of the final tab (level 3 FTE / As-is–To-be comment, decision lines on gap rows, empty comments on
NO GAP / N/A / OUT OF SCOPE rows, system column, to-be split vs decision line, rating vs wording), forbidden text,
terminology, spelling, formulas, yellow cells, contradictions between the two entities, cover glossary.
Findings are reported, never fixed – rebuild from the content file instead.
"""
import argparse
import re
from collections import Counter, defaultdict

import openpyxl
from openpyxl.utils import get_column_letter as L

from fitgap_layout import BLOCK_FIELDS, RATINGS, TemplateLayout, normalize_rating

GAP_RATINGS = {'HARMONIZATION GAP', 'STANDARDIZATION GAP', 'BOTH GAPS'}
NO_COMMENT = {'NO GAP', 'N/A', 'OUT OF SCOPE'}
DECISION_OPENERS = ('Agreement to transfer', 'Agreement in general to transfer', 'Agreement:', 'Standardization gap:',
                    'Deviation:', 'Deviation today;', 'Deviation today:', 'Action LE:', 'Action ', 'Clarification needed:',
                    'No FTE impact', 'Deviation;')
FTE_OPENERS = ('Agreed to transfer', 'Deviation;', 'Deviation:', 'Hours to be defined', 'No FTE impact', 'N/A;', 'N/A ',
               'Negligible', 'Not applicable', 'Out of scope')
FORBIDDEN = [r'\bsession\s+\d', r'\bsessions?\s+(on|of)\s+\d', r'tuesday call', r'live[- ]note', r'open\s*[–-]\s*decision pending',
             r'\bMR NOTES\b', r'reclassified in the session', r'\bin the session\b', r'\bduring the session\b', r'\bas discussed\b',
             r'\bchange[- ]log\b', r'\bKorea\b', r'\bKR\b', r'\[[^\]]{2,}\]', r'\bTODO\b', r'\bxxx\b', r'\?\?', r'\blorem\b']
TERMS = [(r'\bBlackline\b', 'BlackLine'), (r'\bblackline\b', 'BlackLine'), (r'\bSnow\b', 'SNOW'), (r'\bServiceNow request form\b', 'SNOW form'),
         (r'\bemail\b', 'e-mail'), (r'\bEmail\b', 'E-mail'), (r'\bF ?Coloco\b', 'FCLOCO'), (r'\bFcloco\b', 'FCLOCO'),
         (r'\badhoc\b', 'ad hoc'), (r'\bad-hoc\b', 'ad hoc'), (r'\bSCC\b', 'SSC'), (r'\btbd\b', 'TBD'), (r'\bF 101\b', 'F101'),
         (r'\bIFRS COE\b', 'IFRS CoE'), (r'\bWD 1\b', 'WD1'), (r'\bWD 5\b', 'WD5'), (r'\bSSC/GBS\b', 'SSC / GBS')]
TYPOS = ['recieve', 'seperate', 'informaiton', 'orginal', 'occured', 'occurence', 'definately', 'accomodate', 'untill',
         'wich', 'teh ', 'adress', 'transfered', 'reconcilation', 'reconcilliation', 'aquire', 'calender', 'neccessary',
         'recurrent ', 'mantain', 'maintainance', 'standardisation gap', 'harmonisation gap', 'proces ', 'suppliar', 'invoce']


class Report:
    def __init__(self):
        self.sections = []

    def section(self, title, items, level='info', limit=60):
        self.sections.append((title, level, list(items), limit))

    def render(self):
        out = []
        high = sum(1 for t, lv, it, _ in self.sections if lv == 'error' and it)
        out.append(f'# QA report\n\n{high} check(s) with errors, {sum(len(i) for _, _, i, _ in self.sections)} findings in total\n')
        for title, lv, items, limit in self.sections:
            mark = {'error': '❌', 'warn': '⚠️', 'info': 'ℹ️'}[lv] if items else '✅'
            out.append(f'\n## {mark} {title} ({len(items)})\n')
            for x in items[:limit]:
                out.append(f'- {x}')
            if len(items) > limit:
                out.append(f'- … {len(items) - limit} more')
        return '\n'.join(out) + '\n'


def cells_of(ws, lay, codes):
    """Per data row: dict(pid, level, name, row, per entity: fields) for one tab."""
    out = []
    seen = Counter()
    for (r, lv, pid, name, desc) in lay.rows():
        key = pid or f'row:{r}'
        if pid:
            seen[pid] += 1
            if seen[pid] > 1:
                key = f'{pid}#{seen[pid]}'
        per = {}
        for b, code in enumerate(codes):
            cols = lay.blocks[b]['cols']
            vals = {}
            for k in BLOCK_FIELDS:
                c = cols.get(k)
                cell = ws.cell(r, c) if c else None
                v = cell.value if cell else None
                vals[k] = v.strip() if isinstance(v, str) else v
                vals[k + '_cell'] = cell
            per[code] = vals
        out.append({'key': key, 'pid': pid, 'level': lv, 'name': name, 'row': r, 'per': per,
                    'std': lay.std_values(r)})
    return out


def qa(path, capture=None, people=(), out=None):
    rep = Report()
    wb = openpyxl.load_workbook(path)
    names = wb.sheetnames
    fin_name = next((n for n in names if 'fit-gap' in n.lower() and not n.lower().endswith('_vws')), None)
    vws_name = next((n for n in names if n.lower().endswith('_vws')), None)
    # ---- structure
    s = []
    if not fin_name: s.append('final fit-gap tab missing')
    if not vws_name: s.append('_vWS tab missing')
    if fin_name and wb.active.title != fin_name: s.append(f'active sheet is {wb.active.title!r}, not the final tab')
    for must in ('Cover', '>>'):
        if must not in names: s.append(f'sheet {must!r} missing')
    if not any(n.startswith('Process Flows') for n in names): s.append('no Process Flows sheet')
    if fin_name and vws_name and names.index(vws_name) != names.index(fin_name) - 1: s.append('_vWS tab is not directly before the final tab')
    fin = wb[fin_name]; vws = wb[vws_name] if vws_name else None
    lay = TemplateLayout(fin)
    codes = [(re.search(r'\((\w+)\)\s*$', b['label']) or re.search(r'(\w+)$', b['label'] or 'x')).group(1) for b in lay.blocks]
    if len(lay.blocks) != 2: s.append(f'{len(lay.blocks)} entity block(s), expected 2')
    for b in lay.blocks:
        missing = [k for k in BLOCK_FIELDS if not b['cols'].get(k)]
        if missing: s.append(f"block {b['label']!r}: columns missing {missing}")
        cols = [b['cols'][k] for k in BLOCK_FIELDS if b['cols'].get(k)]
        if cols != list(range(cols[0], cols[0] + len(cols))): s.append(f"block {b['label']!r}: columns not contiguous")
    if fin.freeze_panes != f'{L(lay.std["std_le"])}{lay.hdr + 1}': s.append(f'freeze panes {fin.freeze_panes}')
    if not fin.auto_filter.ref or not fin.auto_filter.ref.startswith(f'A{lay.hdr}:'): s.append(f'autofilter {fin.auto_filter.ref}')
    cf_cols = {str(k.sqref).split(':')[0].rstrip('0123456789') for k in fin.conditional_formatting._cf_rules}
    for b in lay.blocks:
        if L(b['cols']['gap']) not in cf_cols: s.append(f"no conditional formatting on GAP column {L(b['cols']['gap'])}")
    dv_sq = ' '.join(str(d.sqref) for d in fin.data_validations.dataValidation)
    for b in lay.blocks:
        for k in ('gap', 'tb_local', 'tb_other', 'tb_gbs'):
            if L(b['cols'][k]) + str(lay.hdr + 1) not in dv_sq: s.append(f'no dropdown on {L(b["cols"][k])}')
    rep.section('Structure', s, 'error')
    # ---- tabs identical except texts
    if vws:
        lay_v = TemplateLayout(vws)
        d = []
        if lay_v.hdr != lay.hdr or lay_v.last_row != lay.last_row: d.append('row range differs between tabs')
        for r in range(1, lay.hdr + 1):
            for c in range(1, fin.max_column + 1):
                if (fin.cell(r, c).value or None) != (vws.cell(r, c).value or None): d.append(f'header differs at {fin.cell(r, c).coordinate}')
        for (r, lv, pid, name, desc) in lay.rows():
            if (vws.cell(r, 3).value, vws.cell(r, 4).value) != (pid, name): d.append(f'row {r}: taxonomy differs between tabs')
        rep.section('vWS vs final tab (header and taxonomy identical)', d[:40], 'error')
    rows = cells_of(fin, lay, codes)
    # ---- fonts, fills, borders
    f = []
    fonts = Counter(); fills_on_rating = []; missing_border = 0; lvl_fill = []
    for row in rows:
        r, lv = row['row'], row['level']
        for c in range(2, fin.max_column + 1):
            cell = fin.cell(r, c)
            if cell.value not in (None, ''):
                fonts[(cell.font.name, cell.font.sz)] += 1
            if lv in (4, 5) and cell.fill.fill_type == 'solid':
                rgb = str(cell.fill.fgColor.rgb)[-6:] if cell.fill.fgColor.type == 'rgb' else 'theme'
                if rgb not in ('FFFF00',):
                    fills_on_rating.append(cell.coordinate)
        for code in codes:
            cell = row['per'][code]['gap_cell']
            if lv in (4, 5) and cell is not None and cell.fill.fill_type == 'solid':
                fills_on_rating.append(cell.coordinate)
    other = {k: v for k, v in fonts.items() if k != ('Verdana', 8.0)}
    if other: f.append(f'fonts other than Verdana 8 in data rows: {other}')
    if fills_on_rating: f.append(f'static fills on level 4/5 cells (first 20): {sorted(set(fills_on_rating))[:20]}')
    rep.section('Formatting (Verdana 8, no static fills on level 4/5)', f, 'warn')
    # ---- ratings
    bad = []
    for row in rows:
        for code in codes:
            g = row['per'][code]['gap']
            if g not in (None, '') and g not in RATINGS:
                bad.append(f"{row['key']} {code}: rating {g!r}")
    rep.section('Ratings outside the six canonical values', bad, 'error')
    # ---- ratings vs capture
    if capture:
        cwb = openpyxl.load_workbook(capture, read_only=False)
        cname = next(n for n in cwb.sheetnames if 'fit' in n.lower() and 'gap' in n.lower())
        clay = TemplateLayout(cwb[cname])
        cmap = {}
        seen = Counter()
        for (r, lv, pid, name, desc) in clay.rows():
            key = pid or f'row:{r}'
            if pid:
                seen[pid] += 1
                if seen[pid] > 1: key = f'{pid}#{seen[pid]}'
            cmap[key] = [clay.block_values(r, b) for b in range(len(codes))]
        changed, missing = [], []
        for row in rows:
            if row['key'] not in cmap:
                missing.append(row['key']); continue
            for b, code in enumerate(codes):
                cg = cmap[row['key']][b].get('gap')
                fg = row['per'][code]['gap']
                if (cg or None) != (fg or None):
                    changed.append(f"{row['key']} {code}: capture {cg!r} → final {fg!r}")
        rep.section('Rows not found in the capture file', missing, 'error')
        rep.section('Ratings changed vs capture (each must be backed by the session)', changed, 'info', 200)
        if len(cmap) != len(rows): rep.section('Row count vs capture', [f'deliverable {len(rows)} rows, capture {len(cmap)} rows'], 'error')
    # ---- content rules on the final tab
    rules = []
    tb = lambda v, k: v.get(k) not in (None, '')
    for row in rows:
        lv, key = row['level'], row['key']
        for code in codes:
            v = row['per'][code]
            g, cmt, fte = v['gap'], v['cmt'], v['fte']
            if g == 'OUT OF SCOPE' or (lv in (1, 2) and not g):
                continue
            if lv == 3:
                if not fte: rules.append(f'{key} {code}: level-3 FTE cell empty')
                elif not (fte.startswith(FTE_OPENERS) or re.match(r'^[A-Z][^;\n]{2,70}; ', fte)): rules.append(f'{key} {code}: level-3 FTE wording {fte[:60]!r} (expected "<qualifier>; …" France style)')
                if not cmt: rules.append(f'{key} {code}: level-3 comment empty')
                elif not (cmt.startswith('As-is:') and '\nTo-be:' in cmt): rules.append(f'{key} {code}: level-3 comment not "As-is: …\\nTo-be: …"')
            elif lv in (4, 5):
                if g in NO_COMMENT and cmt: rules.append(f'{key} {code}: {g} row has a comment')
                if g in GAP_RATINGS:
                    if not cmt: rules.append(f'{key} {code}: {g} row without decision line')
                    else:
                        if '\n' in cmt: rules.append(f'{key} {code}: decision line has several lines')
                        if not cmt.startswith(DECISION_OPENERS): rules.append(f'{key} {code}: decision line opener {cmt[:50]!r}')
                        low = cmt.lower()
                        if g == 'STANDARDIZATION GAP' and 'agreement to transfer' in low: rules.append(f'{key} {code}: "Agreement to transfer" on a STANDARDIZATION row')
                        if g == 'HARMONIZATION GAP' and 'standardization gap' in low: rules.append(f'{key} {code}: "Standardization gap" text on a HARMONIZATION row')
                        if low.startswith('agreement to transfer') and not tb(v, 'tb_gbs'): rules.append(f'{key} {code}: "Agreement to transfer" but no x under GBS')
                        if low.startswith('deviation') and 'stays local' in low and not tb(v, 'tb_local'): rules.append(f'{key} {code}: "stays local" but no x under Local')
                    if not v['sys']: rules.append(f'{key} {code}: system column empty on a rated row')
                if g in ('NO GAP',) and not v['sys'] and (v['as_le'] or v['as_gbs']): rules.append(f'{key} {code}: NO GAP row without system')
                if g and g not in NO_COMMENT and not (v['as_le'] or v['as_gbs']): rules.append(f'{key} {code}: as-is split empty on a rated row')
                if not (tb(v, 'tb_local') or tb(v, 'tb_other') or tb(v, 'tb_gbs')) and g in GAP_RATINGS: rules.append(f'{key} {code}: to-be split empty on a gap row')
    rep.section('Content rules (final tab)', rules, 'error', 400)
    # ---- forbidden text, terminology, spelling – both tabs
    forb, terms, typos, tmpl_typos = [], [], [], []
    block_cols = {b['cols'][k] for b in lay.blocks for k in BLOCK_FIELDS if b['cols'].get(k)}
    for ws, tag in ((fin, 'final'), (vws, 'vWS')):
        if ws is None: continue
        for row in ws.iter_rows(min_row=lay.hdr + 1):
            for cell in row:
                v = cell.value
                if not isinstance(v, str): continue
                low = v.lower()
                if cell.column not in block_cols:
                    if tag == 'final':
                        for t in TYPOS:
                            if t in low: tmpl_typos.append(f'{cell.coordinate}: {t.strip()!r} in {v[:60]!r}')
                    continue
                for pat in FORBIDDEN:
                    if re.search(pat, v, re.I): forb.append(f'{tag} {cell.coordinate}: /{pat}/ in {v[:70]!r}')
                if tag == 'final':
                    for pat, good in TERMS:
                        if re.search(pat, v): terms.append(f'{cell.coordinate}: use {good!r} in {v[:60]!r}')
                    for t in TYPOS:
                        if t in low: typos.append(f'{cell.coordinate}: {t.strip()!r} in {v[:60]!r}')
    rep.section('Forbidden text in the entity blocks (workshop / change-log language, placeholders, other countries)', forb, 'error', 100)
    rep.section('Terminology (final tab, entity blocks)', terms, 'warn', 100)
    rep.section('Spelling (final tab, entity blocks)', typos, 'warn', 100)
    rep.section('Spelling in template columns (taxonomy / process information – template text, not session content)', tmpl_typos, 'info', 40)
    # ---- formulas
    forms = [f'{tag} {c.coordinate}: {c.value[:40]}' for ws, tag in ((fin, 'final'), (vws, 'vWS')) if ws is not None
             for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith('=')]
    rep.section('Formulas in the fit-gap tabs (template leftovers in the trailing columns are expected; none in the entity blocks)',
                [x for x in forms if not any(L(b['cols'][k]) == re.match(r'[A-Z]+', x.split()[1]).group(0) for b in lay.blocks for k in BLOCK_FIELDS)], 'info', 40)
    inblock = [x for x in forms if any(L(b['cols'][k]) == re.match(r'[A-Z]+', x.split()[1]).group(0) for b in lay.blocks for k in BLOCK_FIELDS)]
    rep.section('Formulas inside entity blocks', inblock, 'error')
    # ---- yellow cells
    yel = []
    for ws, tag in ((fin, 'final'), (vws, 'vWS')):
        if ws is None: continue
        for row in ws.iter_rows(min_row=1):
            for c in row:
                if c.fill.fill_type == 'solid' and c.fill.fgColor.type == 'rgb' and str(c.fill.fgColor.rgb).upper().endswith('FFFF00'):
                    txt = str(c.value or '')
                    ok = re.search(r'action|to be confirmed|to be defined|follow-up|pending|clarification', txt, re.I)
                    yel.append(f'{tag} {c.coordinate}{"" if ok else " (no named follow-up in the text!)"}: {txt[:80]!r}')
    rep.section('Yellow cells (only genuinely open topics with a named follow-up)', yel, 'warn', 60)
    # ---- contradictions between the entities
    contra = []
    for row in rows:
        if len(codes) < 2: break
        a, b = row['per'][codes[0]], row['per'][codes[1]]
        if (a['gap'] or None) != (b['gap'] or None): contra.append(f"{row['key']}: rating {a['gap']!r} vs {b['gap']!r}")
        if row['level'] == 3 and (a['cmt'] or '') != (b['cmt'] or ''): contra.append(f"{row['key']}: level-3 comment differs between entities (ok only where the session captured a difference)")
        if row['level'] == 3 and (a['fte'] or '') != (b['fte'] or ''): contra.append(f"{row['key']}: level-3 FTE differs between entities")
    rep.section('Differences between the two entities (review: each must be a captured difference)', contra, 'info', 150)
    # ---- cover
    cov = []
    if 'Cover' in names:
        cw = wb['Cover']
        text = '\n'.join(str(c.value) for row in cw.iter_rows() for c in row if c.value)
        for must in ('Abbreviations', 'People named', 'Sheets', 'Yellow', 'Scope of this version', 'SNOW'):
            if must.lower() not in text.lower(): cov.append(f'cover: section/entry {must!r} missing')
        fonts = {c.font.name for row in cw.iter_rows() for c in row if c.value}
        if fonts - {'Arial'}: cov.append(f'cover fonts other than Arial: {fonts - {"Arial"}}')
        body = '\n'.join(str(c.value) for ws in (fin, vws) if ws is not None for row in ws.iter_rows(min_row=lay.hdr + 1) for c in row if isinstance(c.value, str))
        for p in people:
            if re.search(r'\b' + re.escape(p) + r'\b', body) and p.lower() not in text.lower(): cov.append(f'person {p!r} named in comments but not on the cover')
        # people heuristic: "action <Name>" / "Action <Name>:"
        for m in set(re.findall(r'[Aa]ction ([A-Z][a-z]+):', body)):
            if m not in ('LE', 'SSC', 'GBS', 'PE', 'Ops', 'Local') and m.lower() not in text.lower(): cov.append(f'"action {m}" in comments but {m!r} not on the cover')
    rep.section('Cover', cov, 'error')
    md = rep.render()
    if out:
        open(out, 'w', encoding='utf-8').write(md)
    print(md if not out else md.split('\n')[2])
    return md


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('path'); ap.add_argument('--capture'); ap.add_argument('--people', default=''); ap.add_argument('--out')
    a = ap.parse_args()
    qa(a.path, a.capture, [p.strip() for p in a.people.split(',') if p.strip()], a.out)
